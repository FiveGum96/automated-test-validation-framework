import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# --------------------------------------------------
# Grenzwerte
# --------------------------------------------------

VOLTAGE_MIN = 11.5
VOLTAGE_MAX = 12.5
CURRENT_MAX = 0.50
TEMPERATURE_MAX = 50.0


# --------------------------------------------------
# Messdaten einlesen
# --------------------------------------------------

data = pd.read_csv("data/test_run_001.csv")


# --------------------------------------------------
# Grenzwertverletzungen bestimmen
# --------------------------------------------------

voltage_failures = data[
    (data["voltage_v"] < VOLTAGE_MIN)
    | (data["voltage_v"] > VOLTAGE_MAX)
]

current_failures = data[
    data["current_a"] > CURRENT_MAX
]

temperature_failures = data[
    data["temperature_c"] > TEMPERATURE_MAX
]


# --------------------------------------------------
# PASS / FAIL bestimmen
# --------------------------------------------------

voltage_pass = voltage_failures.empty
current_pass = current_failures.empty
temperature_pass = temperature_failures.empty

overall_pass = (
    voltage_pass
    and current_pass
    and temperature_pass
)


# --------------------------------------------------
# Messwerte bestimmen
# --------------------------------------------------

max_voltage = data["voltage_v"].max()
max_current = data["current_a"].max()
max_temperature = data["temperature_c"].max()


# --------------------------------------------------
# Zeitbereiche bestimmen
# --------------------------------------------------

if voltage_pass:
    voltage_violation = "-"
else:
    voltage_violation = (
        f"{voltage_failures['time_s'].min():.1f} s "
        f"to {voltage_failures['time_s'].max():.1f} s"
    )

if current_pass:
    current_violation = "-"
else:
    current_violation = (
        f"{current_failures['time_s'].min():.1f} s "
        f"to {current_failures['time_s'].max():.1f} s"
    )

if temperature_pass:
    temperature_violation = "-"
else:
    temperature_violation = (
        f"{temperature_failures['time_s'].min():.1f} s "
        f"to {temperature_failures['time_s'].max():.1f} s"
    )


# --------------------------------------------------
# Excel-Datei erstellen
# --------------------------------------------------

workbook = Workbook()

sheet = workbook.active
sheet.title = "Test Summary"


# --------------------------------------------------
# Titel
# --------------------------------------------------

sheet["A1"] = "AUTOMATED TEST VALIDATION REPORT"

sheet["A1"].font = Font(
    bold=True,
    size=16
)


# --------------------------------------------------
# Tabellenüberschriften
# --------------------------------------------------

headers = [
    "Test Case",
    "Requirement",
    "Description",
    "Measured Value",
    "Allowed Limit",
    "Violation Period",
    "Result"
]

for column, header in enumerate(headers, start=1):
    cell = sheet.cell(
        row=3,
        column=column,
        value=header
    )

    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")


# --------------------------------------------------
# Testergebnisse
# --------------------------------------------------

test_results = [
    [
        "TC-001",
        "REQ-001",
        "Voltage Range Test",
        f"{max_voltage:.2f} V",
        "11.50 - 12.50 V",
        voltage_violation,
        "PASS" if voltage_pass else "FAIL"
    ],
    [
        "TC-002",
        "REQ-002",
        "Current Limit Test",
        f"{max_current:.2f} A",
        "<= 0.50 A",
        current_violation,
        "PASS" if current_pass else "FAIL"
    ],
    [
        "TC-003",
        "REQ-003",
        "Temperature Limit Test",
        f"{max_temperature:.2f} °C",
        "<= 50.00 °C",
        temperature_violation,
        "PASS" if temperature_pass else "FAIL"
    ]
]


for row_index, row_data in enumerate(
    test_results,
    start=4
):
    for column_index, value in enumerate(
        row_data,
        start=1
    ):
        sheet.cell(
            row=row_index,
            column=column_index,
            value=value
        )


# --------------------------------------------------
# PASS / FAIL optisch markieren
# --------------------------------------------------

green_fill = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

red_fill = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)


for row in range(4, 7):

    result_cell = sheet.cell(
        row=row,
        column=7
    )

    if result_cell.value == "PASS":
        result_cell.fill = green_fill
    else:
        result_cell.fill = red_fill

    result_cell.font = Font(bold=True)


# --------------------------------------------------
# Gesamtergebnis
# --------------------------------------------------

sheet["A9"] = "Overall Result"
sheet["A9"].font = Font(bold=True)

sheet["B9"] = (
    "PASS"
    if overall_pass
    else "FAIL"
)

sheet["B9"].font = Font(
    bold=True,
    size=12
)

if overall_pass:
    sheet["B9"].fill = green_fill
else:
    sheet["B9"].fill = red_fill


# --------------------------------------------------
# Spaltenbreite automatisch anpassen
# --------------------------------------------------

for column in sheet.columns:

    max_length = 0

    column_letter = get_column_letter(
        column[0].column
    )

    for cell in column:

        if cell.value is not None:

            cell_length = len(
                str(cell.value)
            )

            if cell_length > max_length:
                max_length = cell_length

    sheet.column_dimensions[
        column_letter
    ].width = max_length + 3


# --------------------------------------------------
# Datei speichern
# --------------------------------------------------

output_file = "reports/test_report.xlsx"

workbook.save(output_file)

print(
    f"Excel test report created successfully: "
    f"{output_file}"
)